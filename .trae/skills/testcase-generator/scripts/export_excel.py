"""
测试用例导出工具 - Excel 格式
将 Markdown 测试用例转换为 Excel 文件
"""
import sys
import argparse
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from tc_parser import parse_testcase_markdown


def export_to_excel(testcases: list, output_path: str, module_name: str = "测试用例"):
    """
    将测试用例导出为 Excel 文件
    
    Args:
        testcases: 测试用例字典列表
        output_path: 输出文件路径
        module_name: 工作表名称
    """
    if not testcases:
        print("警告: 没有测试用例可导出")
        return
    
    wb = Workbook()
    ws = wb.active
    ws.title = module_name
    
    # 定义样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    cell_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 写入表头
    headers = list(testcases[0].keys())
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # 写入数据
    for row_num, testcase in enumerate(testcases, 2):
        for col_num, header in enumerate(headers, 1):
            value = testcase.get(header, "")
            cell = ws.cell(row=row_num, column=col_num, value=value)
            cell.alignment = cell_alignment
            cell.border = border
    
    # 调整列宽
    ws.column_dimensions['A'].width = 15  # 用例编号
    ws.column_dimensions['B'].width = 12  # 所属模块
    ws.column_dimensions['C'].width = 30  # 用例标题
    ws.column_dimensions['D'].width = 12  # 用例类型
    ws.column_dimensions['E'].width = 8   # 优先级
    ws.column_dimensions['F'].width = 25  # 前置条件
    ws.column_dimensions['G'].width = 35  # 测试步骤
    ws.column_dimensions['H'].width = 35  # 预期结果
    
    # 冻结首行
    ws.freeze_panes = 'A2'
    
    # 保存文件
    wb.save(output_path)
    print(f"✓ Excel 导出成功: {output_path}")
    print(f"  共 {len(testcases)} 条测试用例")


def main():
    parser = argparse.ArgumentParser(description='将 Markdown 测试用例导出为 Excel')
    parser.add_argument('input', help='输入 Markdown 文件路径')
    parser.add_argument('-o', '--output', help='输出 Excel 文件路径（默认同名 .xlsx）')
    parser.add_argument('-m', '--module', default='测试用例', help='工作表名称')
    
    args = parser.parse_args()
    
    # 读取 Markdown 文件
    input_path = Path(args.input)
    if not input_path.exists():
        print(f"错误: 文件不存在 {input_path}")
        sys.exit(1)
    
    with open(input_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # 解析测试用例
    testcases = parse_testcase_markdown(md_content)
    
    if not testcases:
        print("错误: 未找到测试用例表格")
        sys.exit(1)
    
    # 确定输出路径
    if args.output:
        output_path = Path(args.output)
    else:
        output_path = input_path.with_suffix('.xlsx')
    
    # 导出
    export_to_excel(testcases, str(output_path), args.module)


if __name__ == '__main__':
    main()
